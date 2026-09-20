# -*- coding: utf-8 -*-
"""文档摄入管线（Task 6）：解析 → 版本快照 → 切片 → 索引 → wikilink 落表。

摄入的权威源是 Markdown 正文（``kb_documents.content_md``）：上传解析器
只是「把字节翻译成 MD」的搬运工，向量索引是可重建的派生缓存。一次摄入
的完整链路：

1. **frontmatter 权威**（R6）：MD 内的 ``title`` 与其余键优先于调用参数；
2. **幂等短路**（R4）：同 ``path`` 既有文档、内容哈希一致且已 ``ready``
   时零写入返回——重复摄入同一文件不产生任何副作用；
3. **状态机**：``upsert_document`` 落 ``pending`` → ``processing`` →
   ``ready``；链路任一环异常落 ``failed`` + ``error``（不向上抛，返回
   ``IngestResult(status=failed)``），同内容重摄可重建（失败不钉死）；
4. **切片 → 向量 → 索引**：``seq`` 为文档内序号（三引擎 chunk_id 形态
   统一）；embedding 不可用（``None``）时全空走 BM25-only，摄入不失败；
5. **wikilink 落表**（R8）：出边整替（先删本源旧边再加新边），dst 先按
   原文精确解析、未命中补 ``.md`` 再查，悬挂目标落空串。

fail-soft 边界：上传原文归档（R7）与 space 路由元数据缺失（R14）都不
阻断摄入；写入面异常按「可重触发」处理（落 failed，等待下一次摄入或
T8/T10 的重建任务）。PG 平面不可用（工厂取不到 / 短路探测异常）同样
收敛为 failed 返回，不向上抛。

@author qingfeng
"""

from __future__ import annotations

import io
import logging
import tempfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import yaml

from .chunker import embed_input, split_frontmatter, split_markdown
from .embedding import embed_texts
from .engine import resolve_engine_for
from .file_engine import DEFAULT_KB_DATA_DIR
from .links import extract_wikilinks
from .models import (
    INGEST_FAILED,
    INGEST_PROCESSING,
    INGEST_READY,
    KbDocument,
)
from .pg_store import content_hash, get_ready_kb_pg_store

logger = logging.getLogger(__name__)

#: 一期上传白名单（spec §5.3；pdf/docx 需 ``kb-upload`` extra）
_UPLOAD_SUFFIXES = frozenset(
    {"md", "markdown", "txt", "html", "htm", "pdf", "docx"},
)

#: 上传原文归档根（R7 fail-soft 附加能力；测试可注入覆盖）
_UPLOAD_ROOT = DEFAULT_KB_DATA_DIR / "uploads"


class UnsupportedFormat(ValueError):
    """上传格式不在白名单（一期明确拒绝，不做猜测性解析）。"""


class ParserUnavailable(RuntimeError):
    """格式在路由表内，但解析器依赖未安装（提示安装 ``kb-upload`` extra）。"""


@dataclass(frozen=True)
class IngestResult:
    """一次摄入的产出摘要。

    ``chunk_count`` 是本**次实际写入**的切片数：幂等短路时为 0（没有
    发生任何写入），不代表文档没有切片。
    """

    doc_id: str
    chunk_count: int
    status: str


# ---------------------------------------------------------------------------
# 上传解析（字节 → Markdown 权威源形态）
# ---------------------------------------------------------------------------


def _decode_text(data: bytes) -> str:
    """三级解码（R11）：utf-8-sig（BOM）→ gb18030（存量中文）→ 容错替换。"""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("gb18030")
    except UnicodeDecodeError:
        pass
    return data.decode("utf-8", errors="replace")


def _html_to_md(data: bytes) -> str:
    """HTML → Markdown（html2text，``body_width=0`` 防折行，剔除脚本）。"""
    import html2text

    converter = html2text.HTML2Text()
    # 0 = 关闭自动换行：折行会把中文段落按列宽切开，污染切片
    converter.body_width = 0
    return converter.handle(_decode_text(data))


def _pdf_to_md(data: bytes) -> str:
    """PDF → Markdown（pymupdf4llm，``kb-upload`` extra，延迟 import）。"""
    try:
        import pymupdf4llm
    except ImportError as exc:
        raise ParserUnavailable(
            "PDF 解析需要可选依赖：pip install 'qwenpaw[kb-upload]'",
        ) from exc
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
        handle.write(data)
        pdf_path = handle.name
    try:
        return pymupdf4llm.to_markdown(pdf_path)
    finally:
        Path(pdf_path).unlink(missing_ok=True)


def _docx_to_md(data: bytes) -> str:
    """DOCX → Markdown（python-docx，``kb-upload`` extra，延迟 import）。"""
    try:
        import docx
    except ImportError as exc:
        raise ParserUnavailable(
            "DOCX 解析需要可选依赖：pip install 'qwenpaw[kb-upload]'",
        ) from exc
    document = docx.Document(io.BytesIO(data))
    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    return "\n\n".join(paragraphs)


def parse_upload(filename: str, data: bytes) -> str:
    """把上传字节解析为 Markdown 权威源形态。

    Args:
        filename: 上传文件名（只按扩展名路由，不信任其余部分）。
        data: 原始字节。

    Returns:
        解析后的 Markdown 文本。

    Raises:
        UnsupportedFormat: 扩展名不在白名单。
        ParserUnavailable: pdf/docx 解析器依赖缺失（``kb-upload`` extra）。
    """
    suffix = Path(filename or "").suffix.lower().lstrip(".")
    if suffix not in _UPLOAD_SUFFIXES:
        raise UnsupportedFormat(
            f"暂不支持的上传格式：.{suffix or '(无扩展名)'}",
        )
    if suffix in {"md", "markdown", "txt"}:
        return _decode_text(data)
    if suffix in {"html", "htm"}:
        return _html_to_md(data)
    if suffix == "pdf":
        return _pdf_to_md(data)
    return _docx_to_md(data)


# ---------------------------------------------------------------------------
# 摄入辅助（fail-soft 面与 frontmatter）
# ---------------------------------------------------------------------------


def _frontmatter_of(markdown: str) -> Dict[str, Any]:
    """解析 YAML frontmatter（R6 权威源）；缺失/损坏时返回空 dict。"""
    meta_text, _body = split_frontmatter(markdown)
    if not meta_text.strip():
        return {}
    try:
        loaded = yaml.safe_load(meta_text)
    except Exception:  # pylint: disable=broad-except
        logger.warning("[kb] frontmatter parse failed; ignored")
        return {}
    if not isinstance(loaded, dict):
        return {}
    return loaded


def _parse_object_refs(meta: Dict[str, Any]) -> List[Tuple[str, str]]:
    """解析 frontmatter ``objects`` 键为 (object_type, object_id) 列表。

    支持形态：``list["Project:PROJECT-10001", ...]``、
    ``dict{"Project": ["PROJECT-10001"]}``、逗号分隔字符串。
    无冒号或空段的项静默丢弃（容错不告警：frontmatter 是自由文本）。
    """
    raw = meta.get("objects")
    if raw is None:
        return []
    refs: List[Tuple[str, str]] = []
    if isinstance(raw, str):
        raw = raw.split(",")
    if isinstance(raw, dict):
        for object_type, ids in raw.items():
            id_list = ids if isinstance(ids, (list, tuple)) else [ids]
            refs.extend(
                (str(object_type).strip(), str(value).strip())
                for value in id_list
            )
    elif isinstance(raw, (list, tuple)):
        for item in raw:
            text = str(item).strip()
            if ":" not in text:
                continue
            object_type, _, object_id = text.partition(":")
            refs.append((object_type.strip(), object_id.strip()))
    return [(t, i) for t, i in refs if t and i]


async def _sync_object_links(
    store: Any,
    space_id: str,
    doc_id: str,
    meta: Dict[str, Any],
) -> None:
    """frontmatter ``objects`` → ``kb_object_links``（T4 互引，fail-soft）。

    本体平面不可用 / 对象不存在 / 任一写失败都只记日志，不阻断摄入
    结果（本体互引是附加能力）；relation 默认 ``knowledge_mentions``，
    frontmatter ``objects_relation`` 可覆盖（非法值回落默认）。
    """
    refs = _parse_object_refs(meta)
    if not refs:
        return
    try:
        from ..ontology.models import (
            LINK_RELATION_MENTIONS,
            VALID_LINK_RELATIONS,
        )
        from ..ontology.models import KbObjectLink
        from ..ontology.store import get_ready_ontology_store

        ont_store = await get_ready_ontology_store()
        if ont_store is None:
            logger.info("[kb] object links skipped: ontology plane off")
            return
        relation = str(
            meta.get("objects_relation") or LINK_RELATION_MENTIONS,
        )
        if relation not in VALID_LINK_RELATIONS:
            relation = LINK_RELATION_MENTIONS
        for object_type, object_id in refs:
            if await ont_store.get_object(object_id) is None:
                logger.warning(
                    "[kb] object link skipped (missing object): "
                    "doc=%s object=%s:%s",
                    doc_id,
                    object_type,
                    object_id,
                )
                continue
            link = KbObjectLink(
                id=f"lnk_{uuid.uuid4().hex[:12]}",
                kb_space_id=space_id,
                kb_document_id=doc_id,
                object_type=object_type,
                object_id=object_id,
                relation=relation,
            )
            await ont_store.create_link(link)
    except Exception:  # pylint: disable=broad-except
        logger.warning(
            "[kb] object links sync failed: doc=%s", doc_id, exc_info=True,
        )


def _archive_upload(
    space_id: str,
    doc_id: str,
    source_meta: Dict[str, Any],
    data: bytes,
) -> None:
    """归档上传原文到 ``uploads/<space_id>/<doc_id>.<ext>``（fail-soft）。"""
    try:
        suffix = Path(str(source_meta.get("filename") or "")).suffix.lower()
        directory = Path(_UPLOAD_ROOT) / space_id
        directory.mkdir(parents=True, exist_ok=True)
        (directory / f"{doc_id}{suffix}").write_bytes(data)
    except Exception:  # pylint: disable=broad-except
        logger.warning(
            "[kb] archive upload failed: doc=%s",
            doc_id,
            exc_info=True,
        )


async def _load_space(store: Any, space_id: str) -> Any:
    """读取 space 路由元数据；缺失/异常时返回 ``None``（R14 不阻断）。"""
    try:
        return await store.get_space(space_id)
    except Exception:  # pylint: disable=broad-except
        logger.warning(
            "[kb] load space failed; fall back to defaults: space=%s",
            space_id,
            exc_info=True,
        )
        return None


async def _resolve_links(
    store: Any,
    space_id: str,
    content_md: str,
) -> List[Tuple[str, str, str]]:
    """wikilink → 出边三元组 ``(dst_path, dst_document_id, context)``。

    dst 解析归一化（R8）：先按原文精确查 path；未命中且不以 ``.md`` 结尾
    时补 ``.md`` 再查；仍未命中即悬挂链接（``dst_document_id`` 落空串）。
    """
    resolved: List[Tuple[str, str, str]] = []
    for dst_path, context in extract_wikilinks(content_md):
        dst = await store.get_document_by_path(space_id, dst_path)
        if dst is None and not dst_path.endswith(".md"):
            dst = await store.get_document_by_path(
                space_id,
                f"{dst_path}.md",
            )
        resolved.append(
            (dst_path, dst.id if dst is not None else "", context),
        )
    return resolved


async def _write_chunks(
    *,
    store: Any,
    engine: Any,
    space_id: str,
    doc_id: str,
    title: str,
    content_md: str,
    model: str,
) -> int:
    """写链收口：切片 → 向量 → 索引 → links 落表；返回写入切片数。

    独立成函数是幂等契约的一部分：短路路径不得触达本函数（测试以
    monkeypatch 断言仅首次执行），失败重触发时整链重建
    （``index_document`` 先删后插为底盘）。
    """
    specs = split_markdown(content_md)
    vectors: Optional[List[List[float]]] = None
    if specs:
        # embedding 不可用返回 None：整链走 BM25-only，不降级为失败
        vectors = await embed_texts(
            [embed_input(spec) for spec in specs],
            model=model,
        )
    await engine.index_document(space_id, doc_id, specs, vectors, title=title)
    await store.replace_document_links(
        space_id,
        doc_id,
        await _resolve_links(store, space_id, content_md),
    )
    return len(specs)


# ---------------------------------------------------------------------------
# 摄入主流程
# ---------------------------------------------------------------------------


async def ingest_space_document(
    *,
    space_id: str,
    title: str,
    path: str,
    content_md: str,
    source: str,
    source_meta: Optional[Dict[str, Any]] = None,
    uploaded_from: Optional[bytes] = None,
    store: Any = None,
    engine: Any = None,
) -> IngestResult:
    """摄入一份 Markdown 文档（上传 / URL / 手写统一入口）。

    Args:
        space_id: 目标知识库 ID。
        title: 标题（frontmatter ``title`` 优先于本参数）。
        path: 库内唯一路径（幂等定位键，非空）。
        content_md: Markdown 正文（权威源；空正文拒绝）。
        source: 来源枚举（manual / upload / url）。
        source_meta: 来源元数据（frontmatter 其余键合并覆盖本参数）。
        uploaded_from: 上传原始字节（提供时归档原文，fail-soft）。
        store: KbPgStore（None → PG 工厂；测试注入 fake）。
        engine: 检索引擎（None → 按 space 路由解析；测试注入 fake）。

    Returns:
        :class:`IngestResult`——短路时 ``chunk_count=0``；写链异常
        ``status=failed``（失败态已持久化，可重摄重建）；工厂/短路探测
        故障亦为 ``status=failed``（未建文档，无持久化）。

    Raises:
        ValueError: ``content_md`` 为空白（编程错误，值级拒绝）。
    """
    markdown = content_md or ""
    if not markdown.strip():
        raise ValueError("content_md 为空：没有可摄入的正文")
    if store is None:
        try:
            store = await get_ready_kb_pg_store()
        except Exception:  # pylint: disable=broad-except
            logger.warning("[kb] pg store factory failed", exc_info=True)
            store = None
    if store is None:
        logger.warning(
            "[kb] ingest skipped: PG plane unavailable; path=%s",
            path,
        )
        return IngestResult(doc_id="", chunk_count=0, status=INGEST_FAILED)

    # frontmatter 权威（R6）：title 与其余键覆盖调用参数
    frontmatter = _frontmatter_of(markdown)
    meta: Dict[str, Any] = dict(source_meta or {})
    meta.update(
        {key: value for key, value in frontmatter.items() if key != "title"},
    )
    final_title = str(frontmatter.get("title") or title or "").strip()

    try:
        existing = await store.get_document_by_path(space_id, path)
    except Exception:  # pylint: disable=broad-except
        logger.warning(
            "[kb] short-circuit probe failed: path=%s",
            path,
            exc_info=True,
        )
        return IngestResult(doc_id="", chunk_count=0, status=INGEST_FAILED)
    if (
        existing is not None
        and existing.content_hash == content_hash(markdown)
        and existing.ingest_status == INGEST_READY
    ):
        # 幂等短路（R4）：同 path 同 hash 且 ready —— 零写入返回
        logger.info(
            "[kb] ingest short-circuit: path=%s doc=%s",
            path,
            existing.id,
        )
        return IngestResult(
            doc_id=existing.id,
            chunk_count=0,
            status=INGEST_READY,
        )

    doc_id = (
        existing.id if existing is not None else f"doc_{uuid.uuid4().hex[:12]}"
    )
    if uploaded_from is not None:
        _archive_upload(space_id, doc_id, meta, uploaded_from)

    document = KbDocument(
        id=doc_id,
        space_id=space_id,
        path=path,
        title=final_title,
        content_md=markdown,
        source=source,
        source_meta=meta,
    )
    try:
        await store.upsert_document(document)
        await store.update_ingest_status(doc_id, INGEST_PROCESSING)
        space = await _load_space(store, space_id)
        target_engine = engine or resolve_engine_for(space)
        chunk_count = await _write_chunks(
            store=store,
            engine=target_engine,
            space_id=space_id,
            doc_id=doc_id,
            title=final_title,
            content_md=markdown,
            model=str(getattr(space, "embedding_model", "") or ""),
        )
        await store.update_ingest_status(doc_id, INGEST_READY)
    except Exception as exc:  # pylint: disable=broad-except
        # 写入面异常收敛为 failed 状态（可重触发），不向调用方抛
        error_text = f"{type(exc).__name__}: {exc}"[:500]
        logger.warning(
            "[kb] ingest failed: doc=%s path=%s",
            doc_id,
            path,
            exc_info=True,
        )
        try:
            await store.update_ingest_status(
                doc_id,
                INGEST_FAILED,
                error=error_text,
            )
        except Exception:  # pylint: disable=broad-except
            logger.warning(
                "[kb] persist ingest failure state failed: doc=%s",
                doc_id,
                exc_info=True,
            )
        return IngestResult(
            doc_id=doc_id,
            chunk_count=0,
            status=INGEST_FAILED,
        )
    # T4 知识 ↔ 本体互引（fail-soft：失败不影响摄入结果）
    await _sync_object_links(store, space_id, doc_id, meta)
    return IngestResult(
        doc_id=doc_id,
        chunk_count=chunk_count,
        status=INGEST_READY,
    )


__all__ = [
    "IngestResult",
    "ParserUnavailable",
    "UnsupportedFormat",
    "ingest_space_document",
    "parse_upload",
]
